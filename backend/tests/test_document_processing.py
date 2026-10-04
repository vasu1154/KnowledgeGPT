import os
import fitz
import pytest
from unittest.mock import MagicMock
from docx import Document as DocxDocument
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from app.document_processing.extractor import TextExtractor
from app.document_processing.cleaner import TextCleaner
from app.document_processing.chunker import DocumentChunker
from app.document_processing.processor import DocumentProcessor
from app.ai.embeddings import EmbeddingService
from app.ai.vector_store import VectorStoreService
from app.models.document import Document, DocumentStatus
from app.models.user import User, UserRole


def test_text_extractor_txt(tmp_path):
    txt_file = tmp_path / "sample.txt"
    txt_file.write_text("Hello from plain text document!", encoding="utf-8")

    extractor = TextExtractor()
    pages = extractor.extract(str(txt_file), "txt")
    assert len(pages) == 1
    assert pages[0]["page_number"] == 1
    assert pages[0]["text"] == "Hello from plain text document!"


def test_text_extractor_docx(tmp_path):
    docx_file = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("Paragraph 1 of Word Document.")
    doc.add_paragraph("Paragraph 2 with extra information.")
    doc.save(str(docx_file))

    extractor = TextExtractor()
    pages = extractor.extract(str(docx_file), "docx")
    assert len(pages) == 1
    assert "Paragraph 1" in pages[0]["text"]
    assert "Paragraph 2" in pages[0]["text"]


def test_text_extractor_pdf(tmp_path):
    pdf_file = tmp_path / "sample.pdf"
    doc = fitz.open()

    # Create 2 pages
    page1 = doc.new_page()
    page1.insert_text((50, 72), "Page 1 content here")
    page2 = doc.new_page()
    page2.insert_text((50, 72), "Page 2 content here")
    doc.save(str(pdf_file))
    doc.close()

    extractor = TextExtractor()
    pages = extractor.extract(str(pdf_file), "pdf")
    assert len(pages) == 2
    assert pages[0]["page_number"] == 1
    assert "Page 1" in pages[0]["text"]
    assert pages[1]["page_number"] == 2
    assert "Page 2" in pages[1]["text"]


def test_text_extractor_unsupported():
    extractor = TextExtractor()
    with pytest.raises(ValueError, match="Unsupported file type"):
        extractor.extract("some_file.xyz", "xyz")


def test_text_cleaner():
    cleaner = TextCleaner()
    raw_pages = [
        {
            "page_number": 1,
            "text": "   Line 1 with extra spaces.    \n\n\n\nLine 2 after excessive newlines.   \n   Line 3   ",
        }
    ]
    cleaned = cleaner.clean(raw_pages)
    assert len(cleaned) == 1
    cleaned_text = cleaned[0]["text"]
    assert "\n\n\n" not in cleaned_text
    assert "Line 1 with extra spaces." in cleaned_text
    assert "Line 2 after excessive newlines." in cleaned_text


def test_document_chunker():
    chunker = DocumentChunker(chunk_size=50, chunk_overlap=10)
    pages = [
        {"page_number": 1, "text": "This is a sentence for chunking. " * 3},
        {"page_number": 2, "text": "Second page sentence for testing chunker. " * 2},
    ]

    chunks = chunker.chunk(pages=pages, document_id=42, user_id=7, filename="test.pdf")
    assert len(chunks) > 1

    first_chunk = chunks[0]
    assert "chunk_id" in first_chunk
    assert first_chunk["chunk_id"] == "doc42_chunk0"
    meta = first_chunk["metadata"]
    assert meta["user_id"] == 7
    assert meta["document_id"] == 42
    assert meta["page"] == 1
    assert meta["filename"] == "test.pdf"
    assert meta["chunk_index"] == 0


def test_vector_store_service(tmp_path):
    persist_dir = str(tmp_path / "chroma_test")
    vs = VectorStoreService(persist_directory=persist_dir, collection_name="test_col")

    chunks = [
        {
            "chunk_id": "doc1_chunk0",
            "text": "Python is a versatile programming language.",
            "metadata": {
                "user_id": 10,
                "document_id": 1,
                "chunk_id": "doc1_chunk0",
                "page": 1,
                "filename": "python.txt",
                "chunk_index": 0,
            },
        },
        {
            "chunk_id": "doc2_chunk0",
            "text": "PostgreSQL is an advanced relational database.",
            "metadata": {
                "user_id": 20,
                "document_id": 2,
                "chunk_id": "doc2_chunk0",
                "page": 1,
                "filename": "postgres.txt",
                "chunk_index": 0,
            },
        },
    ]
    # Synthetic 4-dimensional embeddings for test
    embeddings = [
        [0.1, 0.2, 0.3, 0.4],
        [0.9, 0.8, 0.7, 0.6],
    ]

    vs.add_chunks(chunks, embeddings)
    assert vs.get_document_chunk_count(1) == 1
    assert vs.get_document_chunk_count(2) == 1

    # Search with user isolation (user 10 should not see user 20's chunks)
    results = vs.search(query_embedding=[0.1, 0.2, 0.3, 0.4], user_id=10, top_k=2)
    assert len(results) == 1
    assert results[0]["metadata"]["document_id"] == 1

    # Delete doc 1 vectors
    vs.delete_document_vectors(1)
    assert vs.get_document_chunk_count(1) == 0
    assert vs.get_document_chunk_count(2) == 1


@pytest.mark.asyncio
async def test_document_processor_success(tmp_path):
    from app.database import async_session

    # Create dummy text file
    sample_file = tmp_path / "knowledge.txt"
    sample_file.write_text("Artificial Intelligence and Knowledge Retrieval with RAG.", encoding="utf-8")

    async with async_session() as session:
        # Create a test user and document in DB
        user = User(name="Processor User", email="proc_user@example.com", password_hash="hash")
        session.add(user)
        await session.flush()

        doc = Document(
            user_id=user.id,
            filename="proc_doc.txt",
            original_filename="knowledge.txt",
            file_type="txt",
            file_size=len(sample_file.read_bytes()),
            storage_path=str(sample_file),
            status=DocumentStatus.UPLOADED,
        )
        session.add(doc)
        await session.flush()
        doc_id = doc.id

        # Mock embedding service to avoid live external Gemini call in unit test
        mock_embedding_service = MagicMock(spec=EmbeddingService)
        mock_embedding_service.embed_texts.return_value = [[0.1, 0.2, 0.3, 0.4]]

        persist_dir = str(tmp_path / "chroma_proc")
        vector_store = VectorStoreService(persist_directory=persist_dir, collection_name="proc_col")

        processor = DocumentProcessor(
            db=session,
            embedding_service=mock_embedding_service,
            vector_store=vector_store,
        )

        await processor.process_document(doc_id)

        # Check document in DB
        result = await session.execute(select(Document).where(Document.id == doc_id))
        processed_doc = result.scalar_one()

        assert processed_doc.status == DocumentStatus.COMPLETED
        assert processed_doc.page_count == 1
        assert processed_doc.chunk_count >= 1
        assert processed_doc.error_message is None
        assert vector_store.get_document_chunk_count(doc_id) >= 1

        # Clean up
        await session.delete(processed_doc)
        await session.delete(user)
        await session.commit()


@pytest.mark.asyncio
async def test_document_processor_failure(tmp_path):
    from app.database import async_session

    async with async_session() as session:
        user = User(name="Fail User", email="fail_user@example.com", password_hash="hash")
        session.add(user)
        await session.flush()

        doc = Document(
            user_id=user.id,
            filename="missing.pdf",
            original_filename="missing.pdf",
            file_type="pdf",
            file_size=100,
            storage_path=str(tmp_path / "non_existent.pdf"),
            status=DocumentStatus.UPLOADED,
        )
        session.add(doc)
        await session.flush()
        doc_id = doc.id

        processor = DocumentProcessor(db=session)

        with pytest.raises(Exception):
            await processor.process_document(doc_id)

        # Verify status is FAILED and retry_count incremented
        result = await session.execute(select(Document).where(Document.id == doc_id))
        failed_doc = result.scalar_one()

        assert failed_doc.status == DocumentStatus.FAILED
        assert failed_doc.error_message is not None
        assert failed_doc.retry_count == 1

        # Clean up
        await session.delete(failed_doc)
        await session.delete(user)
        await session.commit()
